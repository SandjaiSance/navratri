#!/usr/bin/env python3
"""Analyze Admin_Ritten_2022.xlsx file structure and content."""

try:
    import openpyxl
    from datetime import datetime
    from collections import defaultdict
    
    # Load the workbook
    file_path = r'C:\GithubCopilotWS\DhR\Admin_Ritten_2022.xlsx'
    wb = openpyxl.load_workbook(file_path, data_only=False)
    ws = wb.active
    
    print("=" * 100)
    print("EXCEL FILE ANALYSIS: Admin_Ritten_2022.xlsx")
    print("=" * 100)
    
    # 1. Headers
    print("\n1. SPREADSHEET STRUCTURE - Headers:")
    print("-" * 100)
    
    headers = []
    for col_idx in range(1, ws.max_column + 1):
        header = ws.cell(1, col_idx).value
        if header is not None:
            headers.append(header)
            print(f"  Col {col_idx:2d}: {header}")
    
    print(f"\nTotal columns: {len(headers)}")
    
    # Find key columns
    hours_col = None
    date_col = None
    
    for idx, header in enumerate(headers, 1):
        header_str = str(header).lower()
        if any(x in header_str for x in ['uur', 'hour', 'horas', 'uren']):
            hours_col = idx
            print(f"  --> Hours column found: Column {idx} ({header})")
        if any(x in header_str for x in ['datum', 'date', 'dag']):
            date_col = idx
            print(f"  --> Date column found: Column {idx} ({header})")
    
    # 2. First 5 rows
    print("\n2. FIRST 5 ROWS OF DATA:")
    print("-" * 100)
    
    for row_idx in range(2, min(7, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, min(len(headers) + 1, 9)):
            value = ws.cell(row_idx, col_idx).value
            if value is None:
                row_data.append("[empty]")
            elif isinstance(value, datetime):
                row_data.append(value.strftime('%Y-%m-%d'))
            else:
                row_data.append(str(value)[:20])
        print(f"  Row {row_idx}: {' | '.join(row_data)}")
    
    # 3. Overview
    print("\n3. DATA OVERVIEW:")
    print("-" * 100)
    print(f"  Total rows (including header): {ws.max_row}")
    print(f"  Data rows: {ws.max_row - 1}")
    print(f"  Total columns: {ws.max_column}")
    
    # 4. Hours analysis
    print("\n4. HOURS ANALYSIS:")
    print("-" * 100)
    
    if hours_col:
        total_hours = 0
        record_count = 0
        hours_list = []
        monthly_hours = defaultdict(float)
        monthly_records = defaultdict(int)
        
        for row_idx in range(2, ws.max_row + 1):
            hours_value = ws.cell(row_idx, hours_col).value
            
            if hours_value is not None:
                try:
                    hours = float(hours_value)
                    if hours > 0:
                        total_hours += hours
                        hours_list.append(hours)
                        record_count += 1
                        
                        # Extract month if date available
                        if date_col:
                            date_value = ws.cell(row_idx, date_col).value
                            if isinstance(date_value, datetime):
                                month_key = date_value.month
                                monthly_hours[month_key] += hours
                                monthly_records[month_key] += 1
                except (ValueError, TypeError):
                    pass
        
        print(f"  Total hours (year 2022): {total_hours:.2f}")
        print(f"  Total records with hours: {record_count}")
        
        if record_count > 0:
            print(f"  Average hours per record: {total_hours / record_count:.2f}")
            print(f"  Max hours (single record): {max(hours_list):.2f}")
            print(f"  Min hours (single record): {min(hours_list):.2f}")
        
        # 5. Monthly breakdown
        if monthly_hours:
            print("\n5. MONTHLY BREAKDOWN (2022):")
            print("-" * 100)
            
            month_names = ['', 'January', 'February', 'March', 'April', 'May', 'June',
                          'July', 'August', 'September', 'October', 'November', 'December']
            
            for month in sorted(monthly_hours.keys()):
                month_name = month_names[month] if month < len(month_names) else f"Month {month}"
                hours = monthly_hours[month]
                records = monthly_records[month]
                avg = hours / records if records > 0 else 0
                print(f"  {month_name:12s} (M{month:2d}): {hours:8.2f} hours | {records:3d} records | avg: {avg:6.2f}")
        
        # Summary stats by month
        total_from_months = sum(monthly_hours.values())
        print(f"\n  Total from monthly breakdown: {total_from_months:.2f}")
        if total_from_months != total_hours:
            print(f"  Note: Difference of {total_hours - total_from_months:.2f} (may include records without date)")
    
    print("\n" + "=" * 100)
    print("END OF ANALYSIS")
    print("=" * 100)
    
    wb.close()

except ImportError as e:
    print(f"Error: openpyxl not installed. Install with: pip install openpyxl")
    print(f"Details: {e}")
except Exception as e:
    print(f"Error: {e}")
    import traceback
    traceback.print_exc()
