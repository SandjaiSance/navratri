import openpyxl
from datetime import datetime
from collections import defaultdict

# Load the workbook
wb = openpyxl.load_workbook('C:\\GithubCopilotWS\\DhR\\Admin_Ritten_2022.xlsx')
ws = wb.active

print(f"Sheet name: {ws.title}")
print(f"Dimensions: {ws.dimensions}")

# Check headers
print("\nHeaders (first 15 columns):")
for col in range(1, 16):
    header = ws.cell(1, col).value
    print(f"  Col {col}: {header}")

# Group by month
month_hours = defaultdict(float)
current_row = 2
max_row = ws.max_row

print(f"\nTotal rows: {max_row}")
print("\nAnalyzing data...")

for row in range(2, min(max_row + 1, 100)):  # Check first 98 rows
    date_cell = ws.cell(row, 1).value  # First column is date
    
    # Try to find hours column (need to understand structure)
    if date_cell:
        print(f"Row {row}: Date={date_cell}")
        for col in range(1, 16):
            val = ws.cell(row, col).value
            if val:
                print(f"  Col {col}: {val}")
        print("")

wb.close()
